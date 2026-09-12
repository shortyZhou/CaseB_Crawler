from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path

out_dir = Path('C:/Users/24779/Desktop/AI related/CaseB_Crawler/results')
out_dir.mkdir(parents=True, exist_ok=True)
out_file = out_dir / 'A_result_2.xlsx'

wb = Workbook()
ws = wb.active
ws.title = '汇总'

blue = '1F4E78'
white = 'FFFFFF'
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_sheet(ws, freeze='A2'):
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = border
    for cell in ws[1]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill('solid', fgColor=blue)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            val = '' if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 8), min(max(len(val) + 2, 10), 48))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = ws.dimensions

def add_table(ws, name):
    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)

def add_sheet(title, rows, table_name):
    global ws
    ws = wb.active if title == '汇总' else wb.create_sheet(title)
    if title != '汇总':
        ws.title = title
    for r in rows:
        ws.append(r)
    style_sheet(ws)
    add_table(ws, table_name)

source_url = 'https://mp.weixin.qq.com/s/c2l_RYmeYVaDeR4gO9HU6Q'

add_sheet('汇总', [
    ['数据类型', '萤石', '氢氟酸 / 氟化氢', '备注'],
    ['价格', '97%酸级精粉：3250-3850元/吨；萤石湿粉均价3580元/吨', '无水氟化氢：15100-15300元/吨；市场均价15100元/吨', '来源：微信公众号文章；单位：元/吨'],
    ['电子级价格', '不适用', '电子氢氟酸光伏级：7500-8500元/吨；半导体级：12000-12500元/吨', '9/11与9/10均持平'],
    ['产量/供应', '未给出具体产量；提到多地矿山/选厂停产整改、开工低位', '未给出具体产量；提到北方企业停车检修、行业开工率低位', '为定性描述，不是量化产量'],
    ['消费/需求', '未给出消费量；提到氟化氢复工将改善萤石消耗预期', '未给出消费量；提到后续复工进度影响行情', '为定性描述，不是量化消费量'],
    ['进出口量', '未给出具体量；提到蒙古国进口量稳步增加但同比增量有限', '未提到明确进出口量', '无量化进出口数据'],
    ['市场判断', '短期供应偏紧，矿企挺价惜售，行情高位坚挺', '复工顺利则9月散单商谈价或回调；复工不及预期则稳中偏强', '文章观点整理'],
], 'SummaryTable')

add_sheet('价格数据', [
    ['产品', '9/11价格', '9/10价格', '涨跌', '市场均价', '单位', '来源URL', '备注'],
    ['萤石（97%酸级精粉）', '3250-3850', '未完整显示', '未完整显示', '', '元/吨', source_url, '顶部表格显示该行；文章另称萤石湿粉市场均价3580元/吨'],
    ['萤石湿粉', '', '', '持稳；本周价格重心小幅上移', '3580', '元/吨', source_url, '正文描述'],
    ['无水氟化氢', '15100-15300', '15100-15300', '0', '15100', '元/吨', source_url, '表格行产品名在web_extract中缺失，但正文明确提到无水氟化氢市场均价'],
    ['电子氢氟酸（光伏级）', '7500-8500', '7500-8500', '0', '', '元/吨', source_url, '含氟电子化学品表'],
    ['电子氢氟酸（半导体级）', '12000-12500', '12000-12500', '0', '', '元/吨', source_url, '含氟电子化学品表'],
], 'PriceDataTable')

add_sheet('供应与产量相关', [
    ['品类', '信息类型', '内容', '是否有量化数值', '来源URL'],
    ['萤石', '供应端', '安全检查力度加码，内蒙古部分地区矿山及选厂停产整改', '否', source_url],
    ['萤石', '供应端', '河南、河北地区开工维持低位', '否', source_url],
    ['萤石', '供应端', '南方浙江部分停产矿山尚未复工', '否', source_url],
    ['萤石', '供应端', '常山金石等个别矿企完成整改验收，有望恢复生产，但短期增量有限', '否', source_url],
    ['无水氟化氢 / 氟化氢', '供应端', '北方氟化氢企业停车检修', '否', source_url],
    ['无水氟化氢 / 氟化氢', '开工率', '行业开工率仍处低位', '否', source_url],
    ['无水氟化氢 / 氟化氢', '现货供应', '现货供给阶段性收缩', '否', source_url],
], 'SupplyTable')

add_sheet('消费与需求相关', [
    ['品类', '信息类型', '内容', '是否有量化数值', '来源URL'],
    ['萤石', '下游消耗', '后续氟化氢企业复工对原料萤石的消耗预期向好', '否', source_url],
    ['萤石', '价格支撑', '氟化氢复工预期对萤石价格高位形成支撑', '否', source_url],
    ['无水氟化氢 / 氟化氢', '下游走势', '后期需关注北方企业复工进度', '否', source_url],
    ['无水氟化氢 / 氟化氢', '价格预期', '若顺利恢复生产，9月氟化氢散单商谈价格或有回调；若复工不及预期，短期氟化氢行情将稳中偏强', '否', source_url],
], 'DemandTable')

add_sheet('进出口相关', [
    ['品类', '信息类型', '内容', '是否有量化数值', '来源URL'],
    ['萤石', '进口', '核心进口国蒙古国进口量稳步增加', '否', source_url],
    ['萤石', '进口同比', '同比增量有限', '否', source_url],
    ['萤石', '物流影响', '物流成本上行、运力约束制约实际到货量', '否', source_url],
    ['萤石', '市场影响', '短期国内市场整体货源供应偏紧格局难改', '否', source_url],
    ['氢氟酸 / 氟化氢', '进出口量', '未提到明确进出口量', '否', source_url],
], 'TradeTable')

add_sheet('合规与限制', [
    ['类别', '内容'],
    ['目标文章', source_url],
    ['文章标题', '【氟化工】9月11日市场行情简报！成本端助推，氯化物市场走势偏强'],
    ['来源公众号', '氟务在线'],
    ['robots状态', 'mp.weixin.qq.com/robots.txt 返回 Disallow: /；未继续自动化浏览或深度抓取'],
    ['抓取依据', '仅基于 web_extract 已返回的公开文章文本整理'],
    ['限制1', '顶部表格在 web_extract 输出中部分产品名缺失；无水氟化氢行结合正文识别'],
    ['限制2', '该文主要是市场行情简报，价格信息较完整，产量、消费量、进出口量多为定性描述'],
    ['建议', '若需要完整进出口量数据，建议使用氟务在线官网公开资讯或海关数据源交叉验证'],
], 'ComplianceTable')

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.0' if isinstance(cell.value, float) and not float(cell.value).is_integer() else '#,##0'

wb.save(out_file)
wb2 = load_workbook(out_file, read_only=True, data_only=False)
print(out_file)
print('verified exists:', out_file.exists(), 'size:', out_file.stat().st_size)
for s in wb2.worksheets:
    print(s.title, s.max_row, s.max_column)
