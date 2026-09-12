from pathlib import Path
import sys, subprocess, json
try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

out_dir = Path(r'C:/Users/24779/Desktop/AI related/CaseB_Crawler/results')
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / 'B_result.xlsx'

articles = [
    {
        'date': '2026-09-11',
        'title': '【氟化工】9月11日市场行情简报！成本端助推，氯化物市场走势偏强',
        'account': '氟务在线',
        'url': 'https://mp.weixin.qq.com/s?src=11&timestamp=1789191771&ver=6961&signature=WvkoSeQH4Ysu78nPDRAuR7nH9WozqY*CqWtC-uoUShxudokd3qX7yRl1yVBCUUmyTflorLaYErMde*SMtrk9F0bYz-lJ49aM54Rb*2bFn6cmeu1kjrM9ZaPl11dn4hnQ&new=1',
        'rows': [
            ['萤石', '价格', '萤石（97%酸级精粉）', '3250-3850', '元/吨', '较9/10持平', '表格：氟化工原材料'],
            ['萤石', '价格', '萤石湿粉市场均价', '3580', '元/吨', '当日持稳；本周价格重心小幅上移', '正文'],
            ['氢氟酸', '价格', '无水氟化氢', '15100-15300', '元/吨', '较9/10持平', '表格：氟化工原材料'],
            ['氢氟酸', '价格', '无水氟化氢市场均价', '15100', '元/吨', '正文均价', '正文'],
            ['氢氟酸', '价格', '电子氢氟酸（光伏级）', '7500-8500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['氢氟酸', '价格', '电子氢氟酸（半导体级）', '12000-12500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['萤石', '产量/供应', '矿山/选厂开工', '未给出具体产量数值', '', '内蒙古部分矿山及选厂停产整改；河南、河北开工维持低位；浙江部分停产矿山尚未复工，个别矿企完成整改验收但短期增量有限', '正文'],
            ['氢氟酸', '产量/开工', '北方氟化氢企业', '未给出具体产量数值', '', '停车检修，行业开工率仍处低位，现货供给阶段性收缩', '正文'],
            ['萤石/氢氟酸', '消费/需求', '氟化氢企业复工预期', '未给出具体消费量数值', '', '对原料萤石消耗形成支撑；若顺利恢复生产，9月氟化氢散单商谈价格或回调', '正文'],
            ['萤石', '进出口量', '蒙古国萤石进口', '未给出具体进出口量', '', '进口量稳步增加，但同比增量有限，物流成本、运力约束限制实际到货量', '正文'],
        ],
    },
    {
        'date': '2026-09-10',
        'title': '【氟化工】9月10日市场行情简报！碳酸锂价格下行，氟化锂市场承压',
        'account': '氟务在线',
        'url': 'https://mp.weixin.qq.com/s?src=11&timestamp=1789191742&ver=6961&signature=WvkoSeQH4Ysu78nPDRAuR7nH9WozqY*CqWtC-uoUShxqpmEZMfw9afz9cqvsO9TtaKGB2MpyP916J*dqGlSu7h4-p5U9cmapdCXFmn5XX3m72nV96embLDCbdvgae-N7&new=1',
        'rows': [
            ['萤石', '价格', '萤石（97%酸级精粉）', '3250-3850', '元/吨', '较9/9持平', '表格：氟化工原材料'],
            ['萤石', '价格', '萤石湿粉市场均价', '3580', '元/吨', '持稳', '正文'],
            ['氢氟酸', '价格', '无水氟化氢', '15100-15300', '元/吨', '较9/9持平', '表格：氟化工原材料'],
            ['氢氟酸', '价格', '无水氟化氢市场均价', '15100', '元/吨', '正文均价', '正文'],
            ['氢氟酸', '价格', '电子氢氟酸（光伏级）', '7500-8500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['氢氟酸', '价格', '电子氢氟酸（半导体级）', '12000-12500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['萤石', '产量/供应', '矿山/选厂开工', '未给出具体产量数值', '', '内蒙古部分矿山及选厂仍停产整改；浙江个别矿企完成整改验收即将恢复生产；3万吨及以下小型矿山面临整改、关停压力；市场开工率仍处低位', '正文'],
            ['氢氟酸', '产量/开工', '北方氟化氢企业', '未给出具体产量数值', '', '停车检修，行业开工率仍处低位，现货供给阶段性收缩', '正文'],
            ['萤石/氢氟酸', '消费/需求', '氟化氢行业消耗原料萤石预期', '未给出具体消费量数值', '', '预期向好，支撑萤石价格高位坚挺；需关注北方氟化氢企业复工进度', '正文'],
            ['萤石', '进出口量', '蒙古国萤石进口', '未给出具体进出口量', '', '进口量稳步增加，但物流成本和运力约束制约实际到货量，只能缓解北方货源偏紧', '正文'],
        ],
    },
    {
        'date': '2026-09-09',
        'title': '【氟化工】9月9日市场行情简报！需求不及预期，硫酸跌势难止',
        'account': '氟务在线',
        'url': 'https://mp.weixin.qq.com/s?src=11&timestamp=1789191714&ver=6961&signature=WvkoSeQH4Ysu78nPDRAuR7nH9WozqY*CqWtC-uoUShxx9uqiSGb7wvc6kK1uxpe6Yi5QwiJUjB1I7DWjX5joMY3sVW0polb3Mkdq1GTsJ8wKR40sYdQWlwyGQsQMQp-e&new=1',
        'rows': [
            ['萤石', '价格', '萤石（97%酸级精粉）', '3250-3850', '元/吨', '较9/8持平', '表格：氟化工原材料'],
            ['萤石', '价格', '萤石湿粉市场均价', '3580', '元/吨', '日内上涨5元/吨', '正文'],
            ['氢氟酸', '价格', '无水氟化氢', '15100-15300', '元/吨', '较9/8持平', '表格：氟化工原材料'],
            ['氢氟酸', '价格', '无水氟化氢市场均价', '15100', '元/吨', '正文均价', '正文'],
            ['氢氟酸', '价格', '电子氢氟酸（光伏级）', '7500-8500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['氢氟酸', '价格', '电子氢氟酸（半导体级）', '12000-12500', '元/吨', '持平', '表格：含氟电子化学品'],
            ['萤石', '产量/供应', '矿山/选厂开工', '未给出具体产量数值', '', '内蒙古部分矿山及选厂停产整改；福建部分矿山逐步复工；浙江部分矿山仍未复产；市场开工率维持低位，高品质萤石湿粉紧缺', '正文'],
            ['氢氟酸', '产量/开工', '北方氟化氢企业', '未给出具体产量数值', '', '停车检修，行业开工率仍处低位，现货供给阶段性收缩', '正文'],
            ['萤石/氢氟酸', '消费/需求', '采购/补库', '未给出具体消费量数值', '', '氢氟酸行业消耗原料萤石预期向好，部分企业采购价格小幅上涨；终端跟进有限，多数企业以刚需补库为主，大规模备货未显现', '正文'],
            ['萤石', '进出口量', '蒙古国萤石进口', '未给出具体进出口量', '', '进口量稳步增加，但物流成本上行、运力约束制约实际到货量', '正文'],
        ],
    },
]

wb = Workbook()
ws = wb.active
ws.title = '提取数据'
headers = ['发布日期', '文章标题', '公众号', '产品', '数据类型', '指标', '数值', '单位', '变化/说明', '来源位置', '原文链接']
ws.append(headers)
for art in articles:
    for r in art['rows']:
        ws.append([art['date'], art['title'], art['account'], *r, art['url']])

info = wb.create_sheet('文章信息')
info_headers = ['发布日期', '文章标题', '公众号', '原文链接', '提取状态']
info.append(info_headers)
for art in articles:
    info.append([art['date'], art['title'], art['account'], art['url'], '已获取正文并提取萤石/氢氟酸相关数据'])

summary = wb.create_sheet('摘要')
summary_rows = [
    ['项目', '内容'],
    ['检索工作流', r'C:\Users\24779\Desktop\AI related\CaseB_Crawler\prompts\B_prompt.md'],
    ['检索对象', '公众号“氟务在线”；标题格式“【氟化工】9月X日市场行情简报”'],
    ['日期范围', '2026-09-09 至 2026-09-11'],
    ['提取主题', '萤石或氢氟酸的价格、产量/开工、消费/需求、进出口量'],
    ['总体结论', '三篇文章均含价格数据；均未给出萤石或氢氟酸的具体产量、消费量、进出口量数值，仅提供供应、开工、需求与进口趋势描述。'],
    ['风险提示', '微信 real_url 可能具有时效性；正文价格为企业含税出厂价格，仅供参考。'],
]
for row in summary_rows:
    summary.append(row)

header_fill = PatternFill('solid', fgColor='1F4E78')
header_font = Font(color='FFFFFF', bold=True)
thin = Side(style='thin', color='D9E2F3')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for col in range(1, sheet.max_column + 1):
        max_len = 0
        for cell in sheet[get_column_letter(col)]:
            val = '' if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val), 80))
        sheet.column_dimensions[get_column_letter(col)].width = max(10, min(max_len + 2, 45))
    sheet.auto_filter.ref = sheet.dimensions

for sheet in (ws, info):
    url_col = sheet.max_column if sheet.title == '提取数据' else 4
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, url_col)
        cell.hyperlink = cell.value
        cell.style = 'Hyperlink'

for sheet, name in [(ws, 'ExtractedData'), (info, 'ArticleInfo')]:
    ref = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(tab)

wb._sheets = [summary, ws, info]
wb.save(out_path)

v = load_workbook(out_path, read_only=False, data_only=False)
verify = {s.title: {'rows': s.max_row, 'cols': s.max_column} for s in v.worksheets}
print(json.dumps({'created': str(out_path), 'exists': out_path.exists(), 'size_bytes': out_path.stat().st_size, 'sheets': verify}, ensure_ascii=False, indent=2))
